# import_book.py
#
# One-command book importer.
# Put your PDF or DOCX in data/books/ and run:
#   python import_book.py data/books/book.pdf
#   python import_book.py data/books/book.docx
#
# Step 6: --pipeline and --project-type flags for CLI parity with the
# upload dialog. Default pipeline is "academic" (the book importer).

from __future__ import annotations

import logging
import json
import uuid
from datetime import datetime
from pathlib import Path

from translate_core.doc_parser import DocumentParser
import config

log = logging.getLogger(__name__)

VALID_PROJECT_TYPES = ("book_translation", "article_translation", "festival_programme", "exhibition_catalogue")


def sanitize_lang_pair(pair: str) -> str:
    if not pair:
        return "en->sl"
    cleaned = pair.strip().lower()
    if cleaned == "en-":
        return "en->sl"
    if cleaned == "sl-":
        return "sl->en"
    if "-" in cleaned and "->" not in cleaned:
        cleaned = cleaned.replace("-", "->")
    if "->" not in cleaned:
        return "en->sl"
    return cleaned


def import_book(file_path: str, lang_pair: str = "en->sl",
                pipeline: str = "academic", project_type: str = "book_translation"):
    path = Path(file_path)
    if not path.exists():
        log.error(f"File not found: {path.resolve()}")
        return

    clean_pair = sanitize_lang_pair(lang_pair)

    if pipeline not in ("academic", "simple"):
        log.error(f"Invalid pipeline: {pipeline!r}. Use 'academic' or 'simple'.")
        return
    if project_type not in VALID_PROJECT_TYPES:
        log.error(f"Invalid project_type: {project_type!r}. Use one of {VALID_PROJECT_TYPES}.")
        return

    clean_pair = sanitize_lang_pair(lang_pair)
    suffix = path.suffix.lower()
    log.info(f"\n📖 {path.name}  ({clean_pair}, {pipeline}, {project_type})")

    parser = DocumentParser()
    segments: list[dict] = []
    segments_meta: list = []

    # ── DOCX ────────────────────────────────────────────────────────────
    if suffix == ".docx":
        if pipeline == "academic":
            log.info("   Parsing DOCX as academic (footnote/endnote extraction)…")
            try:
                md = parser.docx_to_markdown(path)
            except Exception as ex:
                log.error(f"\nDOCX academic parse failed: {ex}")
                return
            from translate_core.book_outline import split_paragraphs as _split
            for txt in _split(md, max_chars=config.SEGMENT_MAX_CHARS):
                segments.append({"id": len(segments), "source": txt, "target": "", "status": "pending"})
            segments_meta = []
        else:
            # simple pipeline: paragraph-level extraction preserving docx_para_idx
            import docx as _docx
            log.info("   Parsing DOCX as simple (paragraph-level)…")
            try:
                doc = _docx.Document(str(path))
            except Exception as ex:
                log.error(f"\nDOCX parse failed: {ex}")
                return
            for i, p in enumerate(doc.paragraphs):
                txt = p.text.strip()
                if txt:
                    from translate_core.book_outline import split_paragraphs as _split
                    for chunk in _split(txt, max_chars=config.SEGMENT_MAX_CHARS):
                        segments.append({
                            "id": len(segments),
                            "source": chunk,
                            "target": "",
                            "status": "pending",
                            "docx_para_idx": i,
                        })
            segments_meta = []

    # ── PDF ─────────────────────────────────────────────────────────────
    else:
        log.info("   Parsing PDF with MarkItDown…")
        preprocess = pipeline == "academic"
        try:
            md_text, segments_meta = parser.to_markdown_with_meta(path, preprocess=preprocess)
        except Exception as ex:
            log.error(f"\nParse failed: {ex}")
            log.info("   pip install 'markitdown[pdf]'")
            return

        from translate_core.book_outline import split_paragraphs as _split
        for txt in _split(md_text, max_chars=config.SEGMENT_MAX_CHARS):
            segments.append({"id": len(segments), "source": txt, "target": "", "status": "pending"})

    if not segments:
        log.error("No text extracted from document.")
        return

    # ── Save ───────────────────────────────────────────────────────────
    project_id = str(uuid.uuid4())[:8]
    projects_dir = Path("./data/projects")
    projects_dir.mkdir(parents=True, exist_ok=True)

    (projects_dir / f"{project_id}{path.suffix}").write_bytes(path.read_bytes())

    ws = {
        "id": project_id,
        "filename": path.name,
        "lang_pair": clean_pair,
        "pipeline": pipeline,
        "project_type": project_type,
        "active_index": 0,
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "total": len(segments),
        "done": 0,
        "segments": segments,
    }
    if segments_meta:
        ws["segments_meta"] = segments_meta

    target_json = projects_dir / f"{project_id}.json"
    target_json.write_text(json.dumps(ws, ensure_ascii=False, indent=2), encoding="utf-8")

    # ── Summary ────────────────────────────────────────────────────────
    log.info("\n✅ Done!")
    log.info(f"   Project:  {project_id}")
    log.info(f"   Language: {clean_pair}")
    log.info(f"   Pipeline: {pipeline}")
    log.info(f"   Type:     {project_type}")
    log.info(f"   Segments: {len(segments)}")
    log.info(f"   File:     {target_json}")
    if parser.last_footnote_report:
        r = parser.last_footnote_report
        log.info(f"   Footnotes: {r['defs']} defs, {r['refs']} refs, aligned={r['aligned']}")
    log.info("\n   Open localhost:8080 to translate")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    import argparse

    ap = argparse.ArgumentParser(
        description="Import a PDF or DOCX as a new translation project.",
    )
    ap.add_argument("path", help="PDF or DOCX file")
    ap.add_argument("lang_pair", nargs="?", default="en->sl", help="en->sl (default)")
    ap.add_argument("--pipeline", choices=["academic", "simple"], default="academic",
                    help="Pipeline: academic (footnote restructuring) or simple (preserve formatting). Default: academic")
    ap.add_argument("--project-type", choices=list(VALID_PROJECT_TYPES), default="book_translation",
                    help="KG container project type. Default: book_translation")
    args = ap.parse_args()
    import_book(args.path, args.lang_pair, pipeline=args.pipeline, project_type=args.project_type)