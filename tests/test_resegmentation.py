"""Tests for split_long_paragraph, resegment_pending, and compile_from_template multi-segment."""
from __future__ import annotations
import pytest

from translate_core.book_outline import split_long_paragraph, split_paragraphs, resegment_pending
from translate_core.doc_parser import DocumentParser


# ---------------------------------------------------------------------------
# split_long_paragraph
# ---------------------------------------------------------------------------

class TestSplitLongParagraph:
    def test_short_passthrough(self):
        """Paragraphs under max_chars are returned unchanged."""
        p = "Hello world."
        assert split_long_paragraph(p, 100) == [p]

    def test_sentence_boundary_chunks(self):
        """Long paragraph splits at sentence boundaries, each chunk <= max_chars."""
        text = "First sentence. Second sentence. Third sentence. Fourth sentence."
        chunks = split_long_paragraph(text, 30)
        assert len(chunks) > 1
        for chunk in chunks:
            assert len(chunk) <= 30 or chunk.count(". ") == 0  # single long sentence stays whole

    def test_single_huge_sentence_stays_whole(self):
        """A single sentence longer than max_chars is not split mid-sentence."""
        text = "A" * 900  # no sentence boundary
        chunks = split_long_paragraph(text, 700)
        assert chunks == [text]

    def test_split_paragraphs_integration(self):
        """split_paragraphs uses split_long_paragraph under the hood."""
        text = "First para.\n\n" + "Long. " * 200  # 2nd para is very long
        result = split_paragraphs(text, max_chars=700)
        # All segments should be <= 700 unless single-sentence
        for seg in result:
            assert len(seg) <= 700 or ". " not in seg


# ---------------------------------------------------------------------------
# resegment_pending
# ---------------------------------------------------------------------------

class TestResegmentPending:
    def _make_project(self, segments, active_index=0, done_count=0, meta=None):
        ws = {
            "project_id": "test",
            "filename": "test.pdf",
            "lang_pair": "en->sl",
            "active_index": active_index,
            "saved_at": "2026-01-01T00:00:00",
            "total": len(segments),
            "done": done_count,
            "segments": segments,
        }
        if meta is not None:
            ws["segments_meta"] = meta
        return ws

    def test_done_segments_untouched(self):
        """Done segments are never split."""
        segments = [
            {"id": 0, "source": "A" * 800, "target": "Translation", "status": "done"},
        ]
        ws = self._make_project(segments, done_count=1)
        stats = resegment_pending(ws, 700)
        assert stats["split"] == 0
        assert len(ws["segments"]) == 1

    def test_drafted_segments_untouched(self):
        """Segments with any non-empty target draft are never split."""
        segments = [
            {"id": 0, "source": "A" * 800, "target": "partial", "status": "pending"},
        ]
        ws = self._make_project(segments)
        stats = resegment_pending(ws, 700)
        assert stats["split"] == 0

    def test_pending_empty_target_split(self):
        """Pending segments with empty target and long source get split."""
        long_source = "Sentence one. Sentence two. Sentence three. Sentence four."
        segments = [
            {"id": 0, "source": long_source, "target": "", "status": "pending"},
        ]
        ws = self._make_project(segments)
        stats = resegment_pending(ws, 20)
        assert stats["split"] == 1
        assert stats["after"] > stats["before"]

    def test_ids_renumbered(self):
        """After re-split, segment ids are sequential from 0."""
        segments = [
            {"id": 0, "source": "Short.", "target": "", "status": "pending"},
            {"id": 1, "source": "Sentence one. Sentence two.", "target": "", "status": "pending"},
            {"id": 2, "source": "End.", "target": "Done.", "status": "done"},
        ]
        ws = self._make_project(segments, active_index=0, done_count=1)
        resegment_pending(ws, 20)
        for i, seg in enumerate(ws["segments"]):
            assert seg["id"] == i

    def test_active_index_remapped(self):
        """active_index is remapped to the first child of the formerly active segment."""
        long_source = "Sentence one. Sentence two. Sentence three."
        segments = [
            {"id": 0, "source": "Short.", "target": "", "status": "pending"},
            {"id": 1, "source": long_source, "target": "", "status": "pending"},
        ]
        ws = self._make_project(segments, active_index=1)
        resegment_pending(ws, 20)
        # Segment 1 splits; active_index should point to first child of old index 1
        assert ws["active_index"] >= 1  # after segment 0, first child of old 1

    def test_docx_para_idx_copied_to_children(self):
        """When splitting, docx_para_idx is copied to every child."""
        segments = [
            {"id": 0, "source": "Sentence one. Sentence two.", "target": "", "status": "pending", "docx_para_idx": 5},
        ]
        ws = self._make_project(segments)
        resegment_pending(ws, 20)
        for seg in ws["segments"]:
            assert seg.get("docx_para_idx") == 5

    def test_segments_meta_lockstep_duplication(self):
        """segments_meta is duplicated per child when lengths match."""
        segments = [
            {"id": 0, "source": "Sentence one. Sentence two.", "target": "", "status": "pending"},
        ]
        meta = [{"chapter_index": 3, "page": 42}]
        ws = self._make_project(segments, meta=meta)
        resegment_pending(ws, 20)
        assert len(ws["segments_meta"]) == len(ws["segments"])
        for m in ws["segments_meta"]:
            assert m["chapter_index"] == 3

    def test_segments_meta_mismatch_deleted(self):
        """segments_meta is deleted if length doesn't match segments."""
        segments = [
            {"id": 0, "source": "Short.", "target": "", "status": "pending"},
        ]
        meta = [{"chapter_index": 3}, {"chapter_index": 4}]  # wrong length
        ws = self._make_project(segments, meta=meta)
        resegment_pending(ws, 20)
        assert "segments_meta" not in ws

    def test_noop_when_nothing_qualifies(self):
        """Returns zero-split stats when no segments qualify."""
        segments = [
            {"id": 0, "source": "Short.", "target": "", "status": "pending"},
        ]
        ws = self._make_project(segments)
        stats = resegment_pending(ws, 700)
        assert stats == {"before": 1, "after": 1, "split": 0}

    def test_children_are_pending_with_empty_target(self):
        """All children have status='pending' and target=''."""
        long_source = "First sentence. Second sentence. Third sentence."
        segments = [
            {"id": 0, "source": long_source, "target": "", "status": "pending"},
        ]
        ws = self._make_project(segments)
        resegment_pending(ws, 20)
        for seg in ws["segments"]:
            assert seg["status"] == "pending"
            assert seg["target"] == ""


# ---------------------------------------------------------------------------
# compile_from_template multi-segment join
# ---------------------------------------------------------------------------

class TestCompileFromTemplateMultiSegment:
    @pytest.fixture
    def docx_template(self, tmp_path):
        """Create a minimal DOCX with 3 paragraphs."""
        try:
            import docx
        except ImportError:
            pytest.skip("python-docx not installed")
        doc = docx.Document()
        doc.add_paragraph("Paragraph one original.")
        doc.add_paragraph("Paragraph two original.")
        doc.add_paragraph("Paragraph three original.")
        template_path = tmp_path / "template.docx"
        doc.save(str(template_path))
        return template_path

    def test_multi_segment_join(self, tmp_path, docx_template):
        """When multiple segments share a docx_para_idx, their targets are joined."""
        segments = [
            {"id": 0, "source": "Paragraph one", "target": "Odstavek ena", "status": "done", "docx_para_idx": 0},
            {"id": 1, "source": "Paragraph two part one.", "target": "Dva del ena.", "status": "done", "docx_para_idx": 1},
            {"id": 2, "source": "Paragraph two part two.", "target": "Dva del dva.", "status": "done", "docx_para_idx": 1},
            {"id": 3, "source": "Paragraph three", "target": "Odstavek tri", "status": "done", "docx_para_idx": 2},
        ]
        parser = DocumentParser()
        out_path = tmp_path / "output.docx"
        parser.compile_from_template(docx_template, out_path, segments)

        import docx
        result = docx.Document(str(out_path))
        texts = [p.text for p in result.paragraphs]
        # Paragraph 2 should have both targets joined with space
        assert "Dva del ena." in texts[1]
        assert "Dva del dva." in texts[1]

    def test_all_empty_targets_keeps_source(self, tmp_path, docx_template):
        """When ALL segments for a paragraph have empty target, original is kept."""
        segments = [
            {"id": 0, "source": "Paragraph one", "target": "Prevod ena", "status": "done", "docx_para_idx": 0},
            {"id": 1, "source": "Paragraph two part one.", "target": "", "status": "pending", "docx_para_idx": 1},
            {"id": 2, "source": "Paragraph two part two.", "target": "", "status": "pending", "docx_para_idx": 1},
            {"id": 3, "source": "Paragraph three", "target": "Prevod tri", "status": "done", "docx_para_idx": 2},
        ]
        parser = DocumentParser()
        out_path = tmp_path / "output.docx"
        parser.compile_from_template(docx_template, out_path, segments)

        import docx
        result = docx.Document(str(out_path))
        texts = [p.text for p in result.paragraphs]
        # Paragraph 2 should keep the original because ALL targets are empty
        assert "Paragraph two original." in texts[1]

    def test_partial_empty_targets_falls_back_to_source(self, tmp_path, docx_template):
        """When some segments have targets and others don't, empty ones fall back to source."""
        segments = [
            {"id": 0, "source": "Paragraph one", "target": "Prevod ena", "status": "done", "docx_para_idx": 0},
            {"id": 1, "source": "Paragraph two part one.", "target": "Dva del ena.", "status": "done", "docx_para_idx": 1},
            {"id": 2, "source": "Paragraph two part two.", "target": "", "status": "pending", "docx_para_idx": 1},
            {"id": 3, "source": "Paragraph three", "target": "Prevod tri", "status": "done", "docx_para_idx": 2},
        ]
        parser = DocumentParser()
        out_path = tmp_path / "output.docx"
        parser.compile_from_template(docx_template, out_path, segments)

        import docx
        result = docx.Document(str(out_path))
        texts = [p.text for p in result.paragraphs]
        # Paragraph 2 should have the translated part + source fallback
        assert "Dva del ena." in texts[1]
        assert "Paragraph two part two." in texts[1]