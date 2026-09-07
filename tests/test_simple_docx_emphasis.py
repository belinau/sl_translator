# tests/test_simple_docx_emphasis.py
#
# Tests for the simple-pipeline DOCX emphasis path:
#   1. _docx_paragraph_to_markdown  — DOCX runs → *italic* / **bold** markers
#   2. _parse_docx                  — segments carry emphasis markers
#   3. compile_from_template        — markdown emphasis in target renders
#      as real Word italic/bold runs; plain text falls back to proportional
#      distribution.

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

docx = pytest.importorskip("docx")


# ======================================================================
# _docx_paragraph_to_markdown — DOCX runs → markdown emphasis
# ======================================================================


def _make_paragraph(doc, runs: list[tuple[str, bool, bool]]) -> "docx.paragraph.Paragraph":
    """Helper: build a paragraph with runs as (text, italic, bold)."""
    p = doc.add_paragraph()
    for txt, it, bd in runs:
        run = p.add_run(txt)
        run.italic = it
        run.bold = bd
    return p


class TestDocxParagraphToMarkdown:
    """_docx_paragraph_to_markdown converts runs to markdown emphasis."""

    def test_plain_text_no_markers(self, tmp_path):
        from main import _docx_paragraph_to_markdown
        d = docx.Document()
        p = _make_paragraph(d, [("Plain text.", False, False)])
        assert _docx_paragraph_to_markdown(p) == "Plain text."

    def test_italic_wrapped(self, tmp_path):
        from main import _docx_paragraph_to_markdown
        d = docx.Document()
        p = _make_paragraph(d, [("Regular and ", False, False),
                                 ("italic part", True, False),
                                 (" text.", False, False)])
        result = _docx_paragraph_to_markdown(p)
        assert "*italic part*" in result
        assert "Regular and " in result

    def test_bold_wrapped(self, tmp_path):
        from main import _docx_paragraph_to_markdown
        d = docx.Document()
        p = _make_paragraph(d, [("A ", False, False),
                                 ("bold word", False, True),
                                 (".", False, False)])
        result = _docx_paragraph_to_markdown(p)
        assert "**bold word**" in result

    def test_bold_italic_wrapped(self, tmp_path):
        from main import _docx_paragraph_to_markdown
        d = docx.Document()
        p = _make_paragraph(d, [("both", True, True)])
        result = _docx_paragraph_to_markdown(p)
        assert "***both***" in result

    def test_adjacent_same_flags_coalesced(self, tmp_path):
        """Two adjacent italic runs coalesce into one *span*, not *a**b*."""
        from main import _docx_paragraph_to_markdown
        d = docx.Document()
        p = _make_paragraph(d, [("one", True, False), ("two", True, False)])
        result = _docx_paragraph_to_markdown(p)
        assert result == "*onetwo*"
        assert "**" not in result


# ======================================================================
# _parse_docx — segments carry emphasis markers
# ======================================================================


class TestParseDocxEmphasis:
    """_parse_docx emits segments with markdown emphasis markers."""

    def test_parse_docx_emits_italic_markers(self, tmp_path):
        from main import _parse_docx
        d = docx.Document()
        d.add_paragraph("This has *italic* in it.")  # plain text with literal *
        # Actually we need a real italic run:
        d2 = docx.Document()
        p = d2.add_paragraph()
        p.add_run("This has ")
        r = p.add_run("italic")
        r.italic = True
        p.add_run(" in it.")
        path = tmp_path / "test.docx"
        d2.save(str(path))

        segments = _parse_docx(path)
        assert len(segments) >= 1
        source = segments[0]["source"]
        assert "*italic*" in source
        assert "This has " in source
        assert "docx_para_idx" in segments[0]

    def test_parse_docx_emits_bold_markers(self, tmp_path):
        from main import _parse_docx
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("A ")
        r = p.add_run("bold")
        r.bold = True
        p.add_run(" word.")
        path = tmp_path / "bold.docx"
        d.save(str(path))

        segments = _parse_docx(path)
        assert len(segments) >= 1
        assert "**bold**" in segments[0]["source"]


# ======================================================================
# compile_from_template — markdown emphasis renders in template export
# ======================================================================


class TestCompileFromTemplateEmphasis:
    """compile_from_template renders markdown emphasis as real Word runs."""

    def test_emphasis_target_renders_as_runs(self, tmp_path):
        """Target text with *italic* / **bold** markers → real Word runs,
        no literal asterisks."""
        from translate_core.doc_parser import DocumentParser

        # Build a template DOCX with one paragraph carrying plain runs.
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("Original text here.")
        template = tmp_path / "template.docx"
        d.save(str(template))

        segments = [{
            "id": 0,
            "source": "Original text here.",
            "target": "Prevedeno *kurzivo* in **krepko**.",
            "status": "done",
            "docx_para_idx": 0,
        }]

        out = tmp_path / "out.docx"
        DocumentParser().compile_from_template(template, out, segments)

        result = docx.Document(str(out))
        paras = [par for par in result.paragraphs if par.text.strip()]
        assert len(paras) >= 1
        runs = paras[0].runs
        italic_runs = [r for r in runs if r.italic and not r.bold]
        bold_runs = [r for r in runs if r.bold and not r.italic]
        assert len(italic_runs) >= 1, "Expected italic run from *kurzivo*"
        assert len(bold_runs) >= 1, "Expected bold run from **krepko**"
        # No literal asterisks
        for r in runs:
            assert "*" not in (r.text or ""), f"Literal asterisk in run: {r.text!r}"
        # Content preserved
        italic_text = " ".join(r.text for r in italic_runs)
        assert "kurzivo" in italic_text
        bold_text = " ".join(r.text for r in bold_runs)
        assert "krepko" in bold_text

    def test_plain_text_falls_back_to_proportional(self, tmp_path):
        """Target without emphasis markers uses proportional distribution
        (backward compat)."""
        from translate_core.doc_parser import DocumentParser

        d = docx.Document()
        p = d.add_paragraph()
        r1 = p.add_run("Original ")
        r1.bold = True
        r2 = p.add_run("text.")
        template = tmp_path / "template.docx"
        d.save(str(template))

        segments = [{
            "id": 0,
            "source": "Original text.",
            "target": "Prevedeno besedilo tukaj.",
            "status": "done",
            "docx_para_idx": 0,
        }]

        out = tmp_path / "out.docx"
        DocumentParser().compile_from_template(template, out, segments)

        result = docx.Document(str(out))
        paras = [par for par in result.paragraphs if par.text.strip()]
        # The text should be replaced (no literal asterisks, text present)
        import re as _re
        full_text = _re.sub(r"\s+", " ", " ".join(r.text for r in paras[0].runs))
        assert "Prevedeno besedilo tukaj." in full_text
        # Proportional path keeps the original run count (2 runs: bold + plain)
        assert len(paras[0].runs) == 2

    def test_no_markers_no_emphasis_runs(self, tmp_path):
        """Plain target text produces no italic/bold runs in output."""
        from translate_core.doc_parser import DocumentParser

        d = docx.Document()
        d.add_paragraph("Plain original.")
        template = tmp_path / "template.docx"
        d.save(str(template))

        segments = [{
            "id": 0,
            "source": "Plain original.",
            "target": "Plain translation.",
            "status": "done",
            "docx_para_idx": 0,
        }]

        out = tmp_path / "out.docx"
        DocumentParser().compile_from_template(template, out, segments)

        result = docx.Document(str(out))
        paras = [par for par in result.paragraphs if par.text.strip()]
        for r in paras[0].runs:
            assert not r.italic or r.italic is None
            assert not r.bold or r.bold is None

    def test_template_italic_does_not_leak_to_plain_spans(self, tmp_path):
        """Regression: when the template paragraph's first run is italic
        (e.g. an italicized title) and the target carries *...* markers
        for only part of the text, the plain (unmarked) spans must NOT
        inherit italic from the template's rPr.

        This was the Octavian Esana bug: the entire translation paragraph
        rendered italic because the stripped rPr template still carried
        <w:i>, applied to every plain span.
        """
        from translate_core.doc_parser import DocumentParser

        # Template mimics the real Esana DOCX: run 0 italic (title), run 1
        # plain (description).
        d = docx.Document()
        p = d.add_paragraph()
        r0 = p.add_run("Contemporary Artistic Revolutions")
        r0.italic = True
        p.add_run(". A project historicising the infrastructure.")
        template = tmp_path / "template.docx"
        d.save(str(template))

        segments = [{
            "id": 0,
            "source": "Contemporary Artistic Revolutions. A project historicising the infrastructure.",
            "target": "*Contemporary Artistic Revolutions* (Sodobne umetniške revolucije). Projekt zgodovinjenja infrastrukture.",
            "status": "done",
            "docx_para_idx": 0,
        }]

        out = tmp_path / "out.docx"
        DocumentParser().compile_from_template(template, out, segments)

        result = docx.Document(str(out))
        paras = [par for par in result.paragraphs if par.text.strip()]
        assert len(paras) >= 1
        runs = paras[0].runs

        # No literal asterisks in output
        for r in runs:
            assert "*" not in (r.text or ""), f"Literal asterisk in run: {r.text!r}"

        # The *...* span must be italic
        italic_runs = [r for r in runs if r.italic and not r.bold]
        assert len(italic_runs) >= 1, "Expected the marked title span to be italic"
        assert any("Contemporary Artistic Revolutions" in r.text for r in italic_runs)

        # The plain translation text must NOT be italic — this is the bug.
        plain_runs = [r for r in runs if not r.italic and not r.bold]
        assert len(plain_runs) >= 1, "Expected at least one plain (non-italic) run"
        plain_text = " ".join(r.text for r in plain_runs)
        assert "Sodobne umetniške revolucije" in plain_text, (
            "Translation text should be in a plain run, not an italic one"
        )
        # Explicitly: no plain run should carry italic
        for r in runs:
            if "Sodobne umetniške revolucije" in (r.text or ""):
                assert not r.italic or r.italic is None, (
                    f"Italic leaked into plain translation span: {r.text!r}"
                )


# ======================================================================
# Superscript / soft line break / hyperlink round-trip
# ======================================================================


def _add_hyperlink(paragraph, display: str, url: str) -> None:
    """Append a real external hyperlink to *paragraph* (python-docx has no
    high-level API, so build the XML the same way the export does)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hlink = OxmlElement("w:hyperlink")
    hlink.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = display
    r.append(t)
    hlink.append(r)
    paragraph._element.append(hlink)


class TestSuperscriptBreakLink:
    """Simple-pipeline round-trip for ^superscript^, \\n breaks, [text](url)."""

    def test_superscript_captured_and_rendered(self, tmp_path):
        from main import _docx_paragraph_to_markdown
        from translate_core.doc_parser import DocumentParser
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("Made in GT22 je nov poeticni arhiv")
        sup = p.add_run("2")
        sup.font.superscript = True
        p.add_run(", ki ga tvori.")
        template = tmp_path / "t.docx"
        d.save(str(template))

        md = _docx_paragraph_to_markdown(d.paragraphs[0])
        assert "^2^" in md
        assert "arhiv^2^" in md

        segs = [{"id": 0, "source": md,
                 "target": "Made in GT22 is a new poetic archive^2^, formed.",
                 "status": "done", "docx_para_idx": 0}]
        out = tmp_path / "out.docx"
        DocumentParser().compile_from_template(template, out, segs)
        result = docx.Document(str(out))
        para = [p for p in result.paragraphs if p.text.strip()][0]
        sup_runs = [r for r in para.runs
                    if r.font.superscript is True]
        assert len(sup_runs) == 1, [r.text for r in para.runs]
        assert sup_runs[0].text == "2"

    def test_soft_break_rendered(self, tmp_path):
        from translate_core.doc_parser import DocumentParser
        # Template: one paragraph with a <w:br/> between two bold runs.
        d = docx.Document()
        p = d.add_paragraph()
        r1 = p.add_run("LINE ONE")
        r1.bold = True
        from docx.oxml import OxmlElement
        br_run = p.add_run()
        br_run._r.append(OxmlElement("w:br"))
        r2 = p.add_run("LINE TWO")
        r2.bold = True
        template = tmp_path / "t.docx"
        d.save(str(template))

        segs = [{"id": 0, "source": "**LINE ONE**\n**LINE TWO**",
                 "target": "**FIRST LINE**\n**SECOND LINE**",
                 "status": "done", "docx_para_idx": 0}]
        out = tmp_path / "out.docx"
        DocumentParser().compile_from_template(template, out, segs)
        result = docx.Document(str(out))
        para = [p for p in result.paragraphs if p.text.strip()][0]
        # A <w:br/> must be present in the paragraph XML.
        from docx.oxml.ns import qn
        brs = para._element.findall(".//" + qn("w:br"))
        assert len(brs) == 1, "soft line break not rendered"
        # Two bold runs, no literal newline in run text.
        bold_texts = [r.text for r in para.runs if r.bold and not r.italic]
        assert "FIRST LINE" in bold_texts
        assert "SECOND LINE" in bold_texts

    def test_hyperlink_captured_and_rendered(self, tmp_path):
        from main import _docx_paragraph_to_markdown
        from translate_core.doc_parser import DocumentParser
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("See the ")
        _add_hyperlink(p, "Town Hall", "https://en.wikipedia.org/wiki/Ljubljana_Town_Hall")
        p.add_run(" page.")
        template = tmp_path / "t.docx"
        d.save(str(template))

        md = _docx_paragraph_to_markdown(d.paragraphs[0])
        assert "[Town Hall](https://en.wikipedia.org/wiki/Ljubljana_Town_Hall)" in md

        segs = [{"id": 0, "source": md,
                 "target": "Glej [Mestno hiso](https://sl.wikipedia.org/wiki/Mestna_hisa) stran.",
                 "status": "done", "docx_para_idx": 0}]
        out = tmp_path / "out.docx"
        DocumentParser().compile_from_template(template, out, segs)
        result = docx.Document(str(out))
        para = [p for p in result.paragraphs if p.text.strip()][0]
        # The original hyperlink element must be gone (no bleed) and exactly
        # one new hyperlink with the translated URL present.
        from docx.oxml.ns import qn
        hlinks = para._element.findall(qn("w:hyperlink"))
        assert len(hlinks) == 1
        rid = hlinks[0].get(qn("r:id"))
        rel = result.part.rels[rid]
        assert rel.target_ref == "https://sl.wikipedia.org/wiki/Mestna_hisa"
        display = "".join(t.text or "" for t in hlinks[0].iter(qn("w:t")))
        assert display == "Mestno hiso"

if __name__ == "__main__":
    pytest.main([__file__, "-v"])